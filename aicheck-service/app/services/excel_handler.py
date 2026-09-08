# -*- coding: utf-8 -*-
"""
Excel 用例解析（移植自 checker/excel_handler.py，去掉写回能力）
"""
import os
from typing import List, Dict, Any

from openpyxl import load_workbook


class ExcelHandler:
    """Excel 解析（只读）"""

    def __init__(self, file_path: str):
        self.file_path = os.path.abspath(file_path)
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"Excel 文件不存在：{self.file_path}")
        self._workbook = load_workbook(self.file_path, read_only=True, data_only=True)
        self._sheet = self._workbook.active

    def get_headers(self) -> List[str]:
        headers: List[str] = []
        if not self._sheet:
            return headers
        first_row = next(self._sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        for v in first_row:
            headers.append(str(v).strip() if v is not None else "")
        return headers

    def get_test_cases(self) -> List[Dict[str, Any]]:
        if not self._sheet:
            return []
        headers = self.get_headers()
        cases: List[Dict[str, Any]] = []
        for row in self._sheet.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            case: Dict[str, Any] = {}
            for i, v in enumerate(row):
                if i < len(headers):
                    case[headers[i]] = "" if v is None else v
            if not any(str(case.get(h, "")).strip() for h in headers):
                continue
            cases.append(case)
        return cases

    def close(self):
        if self._workbook:
            self._workbook.close()
            self._workbook = None
            self._sheet = None
